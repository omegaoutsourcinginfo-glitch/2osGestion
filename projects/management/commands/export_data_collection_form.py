"""Génère un formulaire Word (.docx) pour que le collègue remplisse les données projet."""

from pathlib import Path

from django.core.management.base import BaseCommand
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor


MAROON = RGBColor(0x7A, 0x1F, 0x1F)
GRAY = RGBColor(0x55, 0x55, 0x55)


def _set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _add_bottom_border(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "AAAAAA")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _blank_line(doc):
    p = doc.add_paragraph()
    run = p.add_run(" ")
    _set_run_font(run, size=11, color=RGBColor(0xFF, 0xFF, 0xFF))
    _add_bottom_border(p)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(2)


def _field(doc, label, lines=1, hint=""):
    p = doc.add_paragraph()
    run = p.add_run(label)
    _set_run_font(run, size=11, bold=True, color=MAROON)
    if hint:
        h = p.add_run(f"  ({hint})")
        _set_run_font(h, size=9, color=GRAY)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(6)
    for _ in range(lines):
        _blank_line(doc)


def _checkbox_row(doc, options):
    p = doc.add_paragraph()
    for i, opt in enumerate(options):
        if i:
            sp = p.add_run("     ")
            _set_run_font(sp, size=11)
        run = p.add_run(f"☐  {opt}")
        _set_run_font(run, size=11)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(2)


def _section(doc, number, title):
    p = doc.add_paragraph()
    run = p.add_run(f"{number}.  {title}")
    _set_run_font(run, size=14, bold=True, color=MAROON)
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    bar = doc.add_paragraph()
    r = bar.add_run("─" * 60)
    _set_run_font(r, size=8, color=MAROON)
    bar.paragraph_format.space_after = Pt(8)
    bar.paragraph_format.space_before = Pt(0)


def _note(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, size=9, color=GRAY)
    p.paragraph_format.space_after = Pt(6)


def _intervenant_block(doc, type_label):
    p = doc.add_paragraph()
    run = p.add_run(f"▸ {type_label}")
    _set_run_font(run, size=11, bold=True)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    _field(doc, "Nom / Cabinet")
    _field(doc, "Contact (tél. / e-mail)")


def build_document():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("2osGestion")
    _set_run_font(r, size=22, bold=True, color=MAROON)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("FICHE DONNÉES PROJET")
    _set_run_font(r, size=14, bold=True)

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = intro.add_run(
        "À remplir par le collègue — une fiche = un projet.\n"
        "Remplir les champs ci-dessous et renvoyer ce fichier Word.\n"
        "Joindre aussi les fichiers (PDF, photos, plans…) avec le même e-mail / dossier."
    )
    _set_run_font(r, size=10, color=GRAY)

    meta = doc.add_paragraph()
    r = meta.add_run(
        "Date : ______________     Rempli par : ____________________________"
    )
    _set_run_font(r, size=10)
    meta.paragraph_format.space_before = Pt(12)
    meta.paragraph_format.space_after = Pt(4)

    # —— 1. Projet ——
    _section(doc, "1", "PROJET")
    _note(doc, "Le code projet est créé automatiquement — ne pas remplir.")
    _field(doc, "Titre du projet")
    _field(doc, "Préfecture / Province")
    _field(doc, "Commune")
    _field(doc, "Consistance du projet", lines=3)
    _field(doc, "Nombre de lots / Unités")
    _field(doc, "Situation & Localisation", lines=2)
    _field(doc, "Accès au site")
    _field(doc, "Servitudes d'urbanisme", lines=2)
    p = doc.add_paragraph()
    run = p.add_run("Statut actuel :")
    _set_run_font(run, size=11, bold=True, color=MAROON)
    _checkbox_row(doc, ["En cours d'instruction", "En cours d'exécution"])

    # —— 2. Terrain ——
    _section(doc, "2", "TERRAIN MÈRE")
    _field(doc, "Titre foncier mère")
    _field(doc, "Suffixe")
    _field(doc, "Superficie totale")
    p = doc.add_paragraph()
    run = p.add_run("Unité :")
    _set_run_font(run, size=11, bold=True, color=MAROON)
    _checkbox_row(doc, ["m²", "ha"])

    # —— 3. Permis ——
    _section(doc, "3", "PERMIS / AUTORISATION")
    p = doc.add_paragraph()
    run = p.add_run("Type de permis :")
    _set_run_font(run, size=11, bold=True, color=MAROON)
    _checkbox_row(doc, ["Permis d'aménager", "Permis de construire", "Autorisation de lotir"])
    _checkbox_row(
        doc,
        ["Autorisation de morceler", "Permis d'habiter", "Certificat de conformité", "Autre"],
    )
    _field(doc, "Si Autre, préciser")
    _field(doc, "Numéro d'autorisation")
    _field(doc, "Date de délivrance", hint="JJ/MM/AAAA")

    # —— 4. Intervenants ——
    _section(doc, "4", "INTERVENANTS")
    _note(doc, "Laisser vide si non concerné.")
    for label in [
        "Maître d'ouvrage (MO)",
        "Architecte",
        "Ingénieur topographe (IGT)",
        "Bureau d'études techniques (BET)",
        "Bureau de contrôle (BC)",
        "Laboratoire",
        "Entreprise",
        "Autre",
    ]:
        _intervenant_block(doc, label)

    # —— 5. Technique ——
    _section(doc, "5", "PROGRAMME TECHNIQUE")
    p = doc.add_paragraph()
    run = p.add_run("Voirie (VRD) :")
    _set_run_font(run, size=11, bold=True, color=MAROON)
    _checkbox_row(doc, ["Oui", "Non"])
    _field(doc, "Observations", lines=2)

    p = doc.add_paragraph()
    run = p.add_run("Fichiers techniques à envoyer avec cette fiche (cocher) :")
    _set_run_font(run, size=11, bold=True, color=MAROON)
    for kind in [
        "Voirie (VRD)",
        "Télécom / Fibre",
        "Réseau de télécom",
        "Aménagement paysager",
        "Espace vert",
        "STEP",
        "Fiche lot / parcelle",
    ]:
        row = doc.add_paragraph()
        r = row.add_run(f"☐  {kind}")
        _set_run_font(r, size=11)
        row.paragraph_format.space_after = Pt(3)

    # —— 6. Photos ——
    _section(doc, "6", "PHOTOS")
    _note(doc, "Envoyer les photos avec cette fiche. Écrire une légende si besoin.")
    for i in range(1, 6):
        _field(doc, f"Photo {i} — légende")

    # —— 7. Financement ——
    _section(doc, "7", "FINANCEMENT")
    _field(doc, "Fonds propres FP (DH)")
    _field(doc, "Crédit moyen terme CMT (DH)")
    _note(doc, "Joindre les PDF de financement avec cette fiche.")

    # —— 8. Planning ——
    _section(doc, "8", "PLANNING D'EXÉCUTION")
    _note(doc, "Joindre les PDF pour chaque phase.")
    for phase in [
        "Études techniques, géotechniques & Permis",
        "Travaux de Viabilisation (VRD & Réseaux)",
        "Réception provisoire & Titres fonciers individuels",
    ]:
        row = doc.add_paragraph()
        r = row.add_run(f"☐  {phase}")
        _set_run_font(r, size=11)
        row.paragraph_format.space_after = Pt(4)

    # —— 9. Annexes ——
    _section(doc, "9", "ANNEXES (PDF)")
    _note(doc, "Joindre chaque pièce PDF avec cette fiche.")
    for kind in [
        "Certificat de propriété (ANCFCC) — Titre Foncier Mère",
        "Plan Cadastral et Calque topographique (IGT)",
        "Plan d'Architecture & Lotissement « Ne Varietur »",
        "Attestations concessionnaires (Eaux, Électricité, Télécom)",
        "Cahier des charges et PV de réception",
    ]:
        row = doc.add_paragraph()
        r = row.add_run(f"☐  {kind}")
        _set_run_font(r, size=11)
        row.paragraph_format.space_after = Pt(4)

    # —— 10. Remarques ——
    _section(doc, "10", "REMARQUES")
    _field(doc, "Notes", lines=4)

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(28)
    r = footer.add_run("Remplir → renvoyer ce Word + les fichiers — 2osGestion")
    _set_run_font(r, size=9, color=GRAY)

    return doc


class Command(BaseCommand):
    help = "Génère le Word de collecte des données pour le collègue."

    def add_arguments(self, parser):
        parser.add_argument("--out", type=str, default="", help="Chemin .docx de sortie")

    def handle(self, *args, **options):
        doc = build_document()
        desktop = Path.home() / "Desktop"
        docs_dir = Path.cwd() / "docs"
        docs_dir.mkdir(parents=True, exist_ok=True)

        name = "Fiche_Donnees_Projet_COLLEGUE.docx"
        targets = [Path(options["out"])] if options["out"] else [desktop / name, docs_dir / name]

        for path in targets:
            path.parent.mkdir(parents=True, exist_ok=True)
            try:
                doc.save(str(path))
                self.stdout.write(self.style.SUCCESS(f"DOCX : {path}"))
            except PermissionError:
                alt = path.with_name(path.stem + "_NEW.docx")
                doc.save(str(alt))
                self.stdout.write(self.style.WARNING(f"Fichier ouvert — écrit : {alt}"))
